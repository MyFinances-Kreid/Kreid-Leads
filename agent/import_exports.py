"""Intègre les fichiers déposés dans Suivi leads/Imports et trace tout dans un journal.

Usage : python3 import_exports.py <classeur_entrée.xlsx> <dossier Imports> <dossier _agent> <classeur_sortie.xlsx>
Sortie standard : le chemin de l'extraction CRM à utiliser (la plus récente).

- Imports/Weendeal : export du portail Weendeal (CSV ou « CSV Codé », séparateur ; ou ,)
- Imports/Finanzen : export « Liste des prospects » Finanzen (CSV)
- Imports/CRM      : crm_extract_AAAA-MM-JJ.txt produit par le favori « Lire le CRM Kreid »

Règles :
- l'agent AJOUTE en bas des onglets Weendeal / Finanzen les leads absents (clé : e-mail pour Weendeal, téléphone pour Finanzen),
  datés à partir du premier lead déjà suivi dans l'onglet ; il ne modifie jamais une ligne existante ;
- téléphone absent de l'export Weendeal : complété par le CRM (même nom + prénom), sinon signalé ;
- chaque passage écrit dans _agent/journal_imports.csv : fichier lu (empreinte), ajouts, leads déjà présents,
  leads disparus de l'export, changements de statut portail, et pour le CRM les étapes qui ont bougé ;
- l'état du dernier import est gardé dans _agent/etat_imports/ pour comparer au passage suivant.
Données personnelles : journal et état restent dans OneDrive (jamais dans Git).
"""
import sys, os, re, csv, glob, hashlib, unicodedata, datetime as dt, shutil
import openpyxl, copy

WB_IN, IMP, AGENT, WB_OUT = sys.argv[1:5]
STAGE = sys.argv[5] if len(sys.argv) > 5 else None   # run.sh : état et journal publiés seulement si tout le passage réussit
STATE = os.path.join(AGENT, 'etat_imports'); os.makedirs(STATE, exist_ok=True)
STATE_OUT = os.path.join(STAGE, 'etat_imports') if STAGE else STATE; os.makedirs(STATE_OUT, exist_ok=True)
JOURNAL = os.path.join(AGENT, 'journal_imports.csv')
JOURNAL_OUT = os.path.join(STAGE, 'journal_ajout.csv') if STAGE else JOURNAL
NOW = dt.datetime.now().strftime('%Y-%m-%d %H:%M')
EVENTS = []
PRIX = {'Weendeal': 35, 'Finanzen': 48}


def ev(source, fichier, evenement, cle='', nom='', prenom='', detail=''):
    EVENTS.append([NOW, source, fichier, evenement, cle, nom, prenom, detail])


def norm(s):
    s = ''.join(c for c in unicodedata.normalize('NFD', str(s or '')) if unicodedata.category(c) != 'Mn')
    return re.sub(r'[^A-Z ]', ' ', s.upper()).split()


def namekey(nom, prenom):
    return ' '.join(norm(nom) + norm(prenom))


def tel9(x):
    t = re.sub(r'\D', '', str(x or ''))
    if t.startswith('33') and len(t) > 10: t = t[2:]
    t = t.lstrip('0')[-9:]
    return t if len(t) == 9 else ''


def sha(path):
    return hashlib.sha1(open(path, 'rb').read()).hexdigest()[:10]


def latest(folder, pattern):
    files = [f for f in glob.glob(os.path.join(IMP, folder, pattern)) if not os.path.basename(f).startswith('~')]
    return max(files, key=lambda f: (os.path.getmtime(f), f)) if files else None


def read_csv(path):
    raw = open(path, 'rb').read()
    for enc in ('utf-8-sig', 'cp1252'):
        try: txt = raw.decode(enc); break
        except UnicodeDecodeError: continue
    first = txt.split('\n', 1)[0]
    delim = ';' if first.count(';') >= first.count(',') else ','
    return list(csv.DictReader(txt.splitlines(), delimiter=delim)), list(csv.reader([first], delimiter=delim))[0]


def parse_date(s):
    s = str(s or '').strip()
    for f in ('%Y-%m-%d %H:%M:%S', '%Y-%m-%d', '%d/%m/%Y %H:%M:%S', '%d/%m/%Y %H:%M', '%d/%m/%Y'):
        try: return dt.datetime.strptime(s[:19] if 'H' in f else s[:10], f)
        except ValueError: continue
    return None


def num(s):
    s = str(s or '').replace('€', '').replace(' ', '').replace(' ', '').replace(',', '.')
    try: return float(s)
    except ValueError: return None


def load_state(name):
    p = os.path.join(STATE, name + '.csv')
    if not os.path.exists(p): return None, None
    rows = list(csv.reader(open(p, encoding='utf-8'), delimiter=';'))
    meta = dict(x.split('=', 1) for x in rows[0][1:]) if rows and rows[0][0] == '#' else {}
    return meta, {r[0]: r[1:] for r in rows[1:] if r}


def save_state(name, meta, recs):
    with open(os.path.join(STATE_OUT, name + '.csv'), 'w', encoding='utf-8', newline='') as f:
        w = csv.writer(f, delimiter=';')
        w.writerow(['#'] + [f'{k}={v}' for k, v in meta.items()])
        for k, v in recs.items(): w.writerow([k] + v)


# ---------------------------------------------------------------- CRM
crm_file = latest('CRM', 'crm_extract_*.txt')
fallback = sorted(glob.glob(os.path.join(AGENT, 'crm_extract_*.txt')))
crm_use = crm_file
if fallback and (not crm_file or os.path.basename(fallback[-1]) > os.path.basename(crm_file)):
    crm_use = fallback[-1]
CRM_ROWS = [l.split(';') for l in open(crm_use, encoding='utf-8').read().strip().split('\n') if l.count(';') == 10]
CRM_BYNAME = {}
for p in CRM_ROWS:
    nk = namekey(p[3], p[4])
    for t in [t for t in p[2].split(',') if t]:
        CRM_BYNAME.setdefault(nk, set()).add(t[-9:])

STEPS = ['réception', 'simulation', 'retour dossier', 'dépôt', 'accord', 'refus']


def crm_state(rows, keys):
    best = {}
    for p in rows:
        for k in [t[-9:] for t in p[2].split(',') if t]:
            if k not in keys: continue
            sc = sum(1 for x in p[5:11] if x)
            if k not in best or sc > best[k][0]: best[k] = (sc, [p[3], p[4]] + [x.strip() for x in p[5:11]])
    return {k: v[1] for k, v in best.items()}


# ---------------------------------------------------------------- classeur
wb = openpyxl.load_workbook(WB_IN)
W, Fz = wb['Weendeal'], wb['Finanzen']


def tab_rows(ws, r0, ncol):
    return [[ws.cell(r, c).value for c in range(1, ncol + 1)] for r in range(r0, ws.max_row + 1) if ws.cell(r, 1).value not in (None, '')]


def append(ws, values):
    last = max(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value not in (None, ''))
    r = last + 1
    for c, v in enumerate(values, 1):
        x = ws.cell(r, c, v); src = ws.cell(last, c)
        if src.has_style: x._style = copy.copy(src._style)
    return r


# ---------------------------------------------------------------- Weendeal
def do_weendeal():
    src = 'Weendeal'; path = latest('Weendeal', '*.csv')
    if not path: ev(src, '', 'AUCUN_FICHIER', detail='Aucun export dans Imports/Weendeal'); return
    fn = os.path.basename(path); h = sha(path)
    meta, prev = load_state(src)
    rows, head = read_csv(path)
    coded = any(c.startswith('/lead/') for c in head)
    g = (lambda r, k: r.get('/lead/' + k, '')) if coded else None
    recs = {}
    for r in rows:
        if coded:
            d = dict(date=g(r, 'created_on'), nom=g(r, 'nom'), prenom=g(r, 'prenom'), email=g(r, 'email'), tel=g(r, 'tel_mobile') or g(r, 'tel_fixe'),
                     dob=g(r, 'dob'), logement=g(r, 'statut'), profession=g(r, 'profession'), cp=g(r, 'cp'), ville=g(r, 'ville'),
                     statut=r.get('Statut', r.get('/statut', '')))
        else:
            d = dict(date=r.get('Date de demande', ''), nom=r.get('Nom', ''), prenom=r.get('Prénom', ''), email=r.get('Email', ''),
                     tel=r.get('Téléphone', r.get('Tel', '')), dob='', logement='', profession='', cp=r.get('CP', ''), ville='',
                     statut=r.get('Statut', ''))
        k = d['email'].strip().lower() or namekey(d['nom'], d['prenom'])
        if k: recs[k] = d
    ev(src, fn, 'FICHIER_LU', detail=f"empreinte {h} · {len(recs)} leads · format {'CSV Codé' if coded else 'CSV simple'}"
       + ('' if coded else ' (sans téléphone : utiliser « CSV Codé »)'))
    if meta and meta.get('sha') == h:
        ev(src, fn, 'FICHIER_INCHANGE', detail=f"identique à l'import du {meta.get('date')}")
    tab = tab_rows(W, 3, 15)
    in_tab = {str(x[7] or '').strip().lower() for x in tab} | {namekey(x[1], x[2]) for x in tab}
    start = min(parse_date(x[0]) or dt.datetime(2100, 1, 1) for x in tab) if tab else dt.datetime(2000, 1, 1)
    added = present = before = 0
    for k, d in sorted(recs.items(), key=lambda kv: kv[1]['date']):
        dd = parse_date(d['date'])
        if k in in_tab or namekey(d['nom'], d['prenom']) in in_tab: present += 1; continue
        if not dd or dd < start: before += 1; continue
        t = tel9(d['tel']); how = 'export'
        if not t:
            c = CRM_BYNAME.get(namekey(d['nom'], d['prenom']), set())
            if len(c) == 1: t = next(iter(c)); how = 'CRM (même nom)'
        r = append(W, [dd.strftime('%Y-%m-%d %H:%M:%S'), d['nom'], d['prenom'], parse_date(d['dob']), d['logement'] or None, d['profession'] or None,
                       int(t) if t else None, d['email'] or None, d['cp'] or None, d['ville'] or None, PRIX[src], 'non', None, 'non', None])
        added += 1
        ev(src, fn, 'AJOUT_ONGLET', t, d['nom'], d['prenom'], f"demande du {dd:%d/%m/%Y} · ligne {r} · téléphone : {how if t else 'INTROUVABLE (à compléter)'}")
        if not t: ev(src, fn, 'TELEPHONE_MANQUANT', '', d['nom'], d['prenom'], 'ni dans l’export ni dans le CRM : lead non rapprochable')
    if prev is not None:
        for k, d in recs.items():
            if k in prev and prev[k][0] != d['statut']:
                ev(src, fn, 'STATUT_PORTAIL', k, d['nom'], d['prenom'], f"{prev[k][0] or '—'} → {d['statut'] or '—'}")
        for k in prev:
            if k not in recs: ev(src, fn, 'DISPARU_EXPORT', k, prev[k][1], prev[k][2], "présent à l'import précédent, absent de celui-ci")
    ev(src, fn, 'SYNTHESE', detail=f"{added} ajouté(s) · {present} déjà suivis · {before} antérieurs au {start:%d/%m/%Y} (non suivis)"
       + ('' if prev is not None else ' · premier import : base de comparaison créée'))
    save_state(src, {'sha': h, 'date': NOW, 'fichier': fn}, {k: [d['statut'], d['nom'], d['prenom']] for k, d in recs.items()})


# ---------------------------------------------------------------- Finanzen
def do_finanzen():
    src = 'Finanzen'; path = latest('Finanzen', '*.csv')
    if not path: ev(src, '', 'AUCUN_FICHIER', detail='Aucun export dans Imports/Finanzen'); return
    fn = os.path.basename(path); h = sha(path)
    meta, prev = load_state(src)
    rows, _ = read_csv(path)
    recs = {}
    for r in rows:
        k = tel9(r.get('Téléphone')) or (r.get('E-Mail', '').strip().lower())
        if k: recs[k] = r
    ds = [d for d in (parse_date(r.get("Date d'achat")) for r in rows) if d]
    ev(src, fn, 'FICHIER_LU', detail=f"empreinte {h} · {len(recs)} leads" + (f" · achats du {min(ds):%d/%m/%Y} au {max(ds):%d/%m/%Y}" if ds else ''))
    if meta and meta.get('sha') == h:
        ev(src, fn, 'FICHIER_INCHANGE', detail=f"identique à l'import du {meta.get('date')}")
    tab = tab_rows(Fz, 2, 17)
    in_tab = {tel9(x[9]) for x in tab} | {str(x[10] or '').strip().lower() for x in tab}
    start = min(x[0] for x in tab if isinstance(x[0], dt.datetime)) if tab else dt.datetime(2000, 1, 1)
    added = present = before = 0
    for k, r in sorted(recs.items(), key=lambda kv: parse_date(kv[1].get("Date d'achat")) or dt.datetime(2000, 1, 1)):
        dd = parse_date(r.get("Date d'achat"))
        if k in in_tab or r.get('E-Mail', '').strip().lower() in in_tab: present += 1; continue
        if not dd or dd < start: before += 1; continue
        t = tel9(r.get('Téléphone'))
        row = append(Fz, [dd, r.get('Produit'), r.get('Civilité'), r.get('Prénom'), r.get('Nom'), (r.get('Adresse') or '').replace(',', '.'),
                          r.get('Code postal'), r.get('Ville'), parse_date(r.get('Date de naissance')), int(t) if t else None, r.get('E-Mail'),
                          r.get('Profession'), num(r.get('Prix (TTC)')) or PRIX[src], 'non', None, 'non', None])
        added += 1
        ev(src, fn, 'AJOUT_ONGLET', t, r.get('Nom'), r.get('Prénom'), f"acheté le {dd:%d/%m/%Y} · ligne {row} · {r.get('Prix (TTC)', '')}")
    if prev is not None:
        for k, r in recs.items():
            if k in prev and prev[k][0] != r.get('Etat', ''):
                ev(src, fn, 'STATUT_PORTAIL', k, r.get('Nom'), r.get('Prénom'), f"{prev[k][0] or '—'} → {r.get('Etat') or '—'}")
        for k in prev:
            if k not in recs: ev(src, fn, 'DISPARU_EXPORT', k, prev[k][1], prev[k][2], "présent à l'import précédent, absent de celui-ci (export limité à une période ?)")
    ev(src, fn, 'SYNTHESE', detail=f"{added} ajouté(s) · {present} déjà suivis · {before} antérieurs au {start:%d/%m/%Y}"
       + ('' if prev is not None else ' · premier import : base de comparaison créée'))
    save_state(src, {'sha': h, 'date': NOW, 'fichier': fn}, {k: [r.get('Etat', ''), r.get('Nom', ''), r.get('Prénom', '')] for k, r in recs.items()})


do_weendeal()
do_finanzen()


# ---------------------------------------------------------------- CRM : évolution des fiches des leads
def lead_keys():
    ks = set()
    for ws, r0, col in ((W, 3, 7), (Fz, 2, 10), (wb['Qwamplify'], 2, 6)) + (((wb['Kreid.fr'], 2, 4),) if 'Kreid.fr' in wb.sheetnames else ()):
        for r in range(r0, ws.max_row + 1):
            t = tel9(ws.cell(r, col).value)
            if t: ks.add(t)
    return ks


KEYS = lead_keys()
fn = os.path.basename(crm_use); h = sha(crm_use)
meta, _ = load_state('CRM')
ev('CRM', fn, 'FICHIER_LU', detail=f"empreinte {h} · {len(CRM_ROWS)} fiches CRM · " + ('dépôt Imports/CRM' if crm_use == crm_file else 'aucun nouveau dépôt : dernière extraction connue'))
prev_path = os.path.join(STATE, 'crm_precedent.txt')
if meta and meta.get('sha') == h:
    ev('CRM', fn, 'FICHIER_INCHANGE', detail=f"déjà intégré le {meta.get('date')} : relancer le favori « Lire le CRM Kreid »")
else:
    if not os.path.exists(prev_path) and len(fallback) and fallback[-1] != crm_use:
        prev_path = fallback[-1]   # 1re fois : base = dernière extraction de l'agent
    if os.path.exists(prev_path):
        old = crm_state([l.split(';') for l in open(prev_path, encoding='utf-8').read().strip().split('\n') if l.count(';') == 10], KEYS)
        new = crm_state(CRM_ROWS, KEYS)
        n_new = n_mv = 0
        for k, v in new.items():
            if k not in old:
                n_new += 1; ev('CRM', fn, 'NOUVELLE_FICHE_LEAD', k, v[0], v[1], f"reçue le {v[2] or '?'}")
                continue
            for i, s in enumerate(STEPS):
                if v[2 + i] and v[2 + i] != old[k][2 + i]:
                    n_mv += 1; ev('CRM', fn, 'ETAPE_CRM', k, v[0], v[1], f"{s} : {old[k][2 + i] or '—'} → {v[2 + i]}")
        gone = [k for k in old if k not in new]
        for k in gone: ev('CRM', fn, 'FICHE_DISPARUE', k, old[k][0], old[k][1], "fiche du lead absente de la nouvelle extraction (faux contact / supprimée ?)")
        ev('CRM', fn, 'SYNTHESE', detail=f"{len(new)} fiches de leads · {n_new} nouvelle(s) · {n_mv} étape(s) franchie(s) · {len(gone)} disparue(s)")
    else:
        ev('CRM', fn, 'SYNTHESE', detail='premier import : base de comparaison créée')
    shutil.copy(crm_use, os.path.join(STATE_OUT, 'crm_precedent.txt'))
    save_state('CRM', {'sha': h, 'date': NOW, 'fichier': fn}, {})

wb.save(WB_OUT)
new_file = not os.path.exists(JOURNAL)
with open(JOURNAL_OUT, 'a', encoding='utf-8', newline='') as f:
    w = csv.writer(f, delimiter=';')
    if new_file and not STAGE: w.writerow(['horodatage', 'source', 'fichier', 'evenement', 'cle', 'nom', 'prenom', 'detail'])
    w.writerows(EVENTS)
sys.stderr.write('\n'.join(';'.join(map(str, e[1:4] + e[7:])) for e in EVENTS if e[3] in ('SYNTHESE', 'FICHIER_LU', 'FICHIER_INCHANGE', 'TELEPHONE_MANQUANT', 'AUCUN_FICHIER')) + '\n')
print(crm_use)
